import contextlib
import csv
import io
import tempfile
import unittest
from pathlib import Path
from datetime import date
from openpyxl import Workbook, load_workbook
from test_reglas import report

class ExportTest(unittest.TestCase):
    def test_inactive_and_vacation_alerts_with_full_export(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            base = root / 'base'
            base.mkdir()
            config = root / 'config'
            config.mkdir()
            report.save_json(config/'recursos.json',[{'documento':'1','nombre':'Ana Perez','correo':'ana@example.com','alias':[],'activo':True,'ausencias':[],'horas_dia':9,'dedicacion':1,'datos_base':{'Celula interna BD':'DB Tech','EVC':'EVC interno'}}])
            report.save_json(config/'sprints.json',[{'sprint':'255','inicio':'2026-10-07','fin':'2026-10-20','activo':True}])
            report.save_json(config/'parametros.json',{'estados_finales':['Closed','Removed','Resolved'],'tipos_carga':['Task'],'tipos_ha':['Habilitador','User Story']})
            report.save_json(config/'tiempos.json',report.DEFAULT_TIMES)
            report.save_json(config/'festivos.json',{'fechas':['2026-10-12']})
            headers = ['ID','Parent','Work Item Type','Title','Assigned To','State','Iteration Level 3','Original Estimate','Story Points','Area Level 4','Area Level 3']
            with (base / 'David Gallo_AZURE DEVOPS2.csv').open('w',encoding='utf-8',newline='') as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerow(['1','','Habilitador','=1+1','Ana Perez <ana@example.com>','Active','Sprint 255','','3','Celula Azure','EVC Azure'])
                writer.writerow(['2','1','Task','Task','Ana Perez <ana@example.com>','Active','Sprint 255','99','8','Celula Azure','EVC Azure'])
            import json
            people = json.loads((config / 'recursos.json').read_text(encoding='utf-8'))
            report.save_json(config / 'festivos.json', {'fechas': ['2026-10-12']})
            people[0]['activo'] = False
            report.save_json(config / 'recursos.json', people)
            output = root / 'report.xlsx'
            with contextlib.redirect_stdout(io.StringIO()): output = report.run(base, config, output, date(2026,10,8))
            self.assertRegex(output.name, r'^cargabilidad_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}\.xlsx$')
            self.assertTrue(output.with_suffix('.html').is_file())
            book = load_workbook(output)
            kinds = [r[0] for r in book['Alertas'].iter_rows(min_row=2,values_only=True)]
            self.assertEqual(kinds.count('RECURSO_NO_DISPONIBLE'), 2)
            self.assertEqual(book['Detalle'].max_row, 3)
            self.assertEqual(book['Detalle']['D2'].data_type, 's')
            self.assertEqual(book['Detalle']['L2'].value, 'Celula Azure')
            self.assertEqual(book['Detalle']['M2'].value, 'EVC Azure')
            self.assertEqual(book['Detalle']['T3'].value, '1')
            self.assertEqual(book['Carga recursos']['H2'].value, 0)
            self.assertEqual(book['Carga recursos']['D2'].value, 11)
            self.assertEqual(book['Carga recursos']['F2'].value, 11)
            self.assertEqual(book['Config festivos']['A2'].value, '2026-10-12')
            book.close()
            people[0]['activo'] = True
            people[0]['vacaciones'] = {'inicio':'2026-10-07','fin':'2026-10-09'}
            report.save_json(config / 'recursos.json', people)
            with contextlib.redirect_stdout(io.StringIO()): output = report.run(base, config, output, date(2026,10,8))
            self.assertRegex(output.name, r'^cargabilidad_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}\.xlsx$')
            self.assertTrue(output.with_suffix('.html').is_file())
            book = load_workbook(output)
            kinds = [r[0] for r in book['Alertas'].iter_rows(min_row=2,values_only=True)]
            self.assertEqual(kinds.count('RECURSO_NO_DISPONIBLE'), 2)
            self.assertEqual(kinds.count('AUSENCIA_EN_SPRINT'), 2)
            self.assertEqual(book['Carga recursos']['H2'].value, 41)
            self.assertTrue(book['Carga recursos'].row_dimensions[3].hidden)
            self.assertEqual(book['Carga recursos']['B3'].value,'Sprint 255')
            self.assertEqual(book['Historico recursos']['B3'].value,'Sprint 255')
            self.assertEqual(book['Metricas sprint']['A3'].value,'Ana Perez')
            organizational=book['Detalle celulas EVC']
            self.assertNotIn('Sprint',[c.value for c in organizational[1]])
            self.assertEqual(organizational['A4'].value,'Sprint 255')
            self.assertEqual(organizational.row_dimensions[4].outlineLevel,2)
            self.assertEqual(organizational['A5'].value,'Ana Perez')
            self.assertEqual(organizational.row_dimensions[5].outlineLevel,3)
            self.assertEqual(organizational['B2'].value,2)

            self.assertNotIn('Base CSV',book.sheetnames)
            self.assertTrue(all(not ws._charts for ws in book))
            self.assertAlmostEqual(book['Carga recursos']['I2'].value, 11/41)
            self.assertEqual(book['Historico recursos']['C2'].value, 11)
            self.assertEqual(book['Historico recursos']['E2'].value, 11)
            self.assertEqual(book['Matriz por estado']['F2'].value, 11)
            self.assertNotIn('ASIGNADAS - tipos de carga',[r[0] for r in book['Matriz por estado'].iter_rows(values_only=True)])
            hour_details=[r for r in book['Tablero'].iter_rows(values_only=True) if r[2]=='Detalle horas por sprint']
            self.assertTrue(hour_details)
            self.assertTrue(all(r[3]=='255' for r in hour_details))

            self.assertEqual(book['Metricas sprint']['C2'].value, 11)
            self.assertEqual(book['Metricas sprint']['H2'].value, 41)
            self.assertEqual(book.sheetnames[0], 'Tablero')
            dashboard=list(book['Tablero'].iter_rows(values_only=True))
            sprint_totals=[row for row in dashboard if row[2]=='Subtotal sprint']
            type_totals=[row for row in dashboard if row[2]=='Tipo de elemento']
            self.assertEqual(sum(row[1] for row in sprint_totals),2)
            self.assertEqual(sum(row[1] for row in type_totals),2)
            self.assertTrue(all(row[3]=='255' for row in sprint_totals+type_totals))

            book.close()

    def test_batch_paths_and_unavailable_open_items_without_hours(self):
        import json
        import subprocess
        with tempfile.TemporaryDirectory(prefix='cargabilidad prueba ') as temporary:
            root = Path(temporary)
            base, config, output = root/'entrada con espacios', root/'config con espacios', root/'salida con espacios'
            base.mkdir(); config.mkdir()
            people = []
            for name in ('Vacaciones', 'Ausencia', 'Inactivo', 'Incapacidad', 'Disponible'):
                person = {'documento':name,'nombre':name,'correo':name.lower()+'@example.com','alias':[],
                    'activo':name!='Inactivo','horas_dia':9,'dedicacion':1,'ausencias':[],
                    'datos_base':{'Celula interna BD':'Equipo'}}
                if name == 'Vacaciones': person['vacaciones'] = {'inicio':'2026-10-07','fin':'2026-10-09'}
                if name == 'Ausencia': person['ausencias'] = [{'inicio':'2026-10-08','fin':'2026-10-08','motivo':'Permiso'}]
                if name == 'Incapacidad': person['incapacidades'] = [{'inicio':'2026-10-08','fin':'2026-10-09'}]
                people.append(person)
            report.save_json(config/'recursos.json',people)
            report.save_json(config/'sprints.json',[{'sprint':'255','inicio':'2026-10-07','fin':'2026-10-20','activo':True}])
            report.save_json(config/'parametros.json',{'estados_finales':['Closed','Removed','Resolved'],'tipos_carga':['Task'],'tipos_ha':[]})
            report.save_json(config/'tiempos.json',report.DEFAULT_TIMES)
            report.save_json(config/'festivos.json',{'fechas':[]})
            administrative_path = Path(__file__).resolve().parents[1]/'config'/'tareas_administrativas.json'
            report.save_json(config/'tareas_administrativas.json',json.loads(administrative_path.read_text(encoding='utf-8')))
            headers=['ID','Parent','Work Item Type','Title','Assigned To','State','Iteration Level 3','Original Estimate','Story Points','Area Level 4','Area Level 3']
            with (base/'datos.csv').open('w',encoding='utf-8',newline='') as f:
                writer=csv.writer(f);writer.writerow(headers)
                for i,person in enumerate(people,1):
                    writer.writerow([str(i),'','Task','Caso sin horas',person['nombre']+' <'+person['correo']+'>','Active','Sprint 255','','','Celula','EVC'])
                writer.writerow(['6','','Task','Creacion de usuario','Vacaciones <vacaciones@example.com>','Closed','Sprint 255','','8','Celula','EVC'])
                writer.writerow(['7','','Task','Ejecutar script','Disponible <disponible@example.com>','Active','Sprint 255','','','Celula','EVC'])
                writer.writerow(['8','','Habilitador','Otorgar permisos','Disponible <disponible@example.com>','Closed','Sprint 255','','5','Celula','EVC'])
                writer.writerow(['9','','Task','Crear usuario y otorgar permisos','Disponible <disponible@example.com>','Active','Sprint 255','','8','Celula','EVC'])
            import os
            if os.name == 'nt':
                batch = Path(__file__).resolve().parents[1]/'ejecutar.bat'
                result = subprocess.run(['cmd','/c',str(batch),str(base),str(output),'2026-10-08',str(config)],capture_output=True,text=True,timeout=60)
                self.assertEqual(result.returncode,0,result.stdout+result.stderr)
            else:
                with contextlib.redirect_stdout(io.StringIO()): report.run(base,config,output/'cargabilidad.xlsx',date(2026,10,8))
            html_paths=list(output.glob('*.html'));self.assertEqual(len(html_paths),1)
            html=html_paths[0].read_text(encoding='utf-8')
            payload=json.loads(html.split('<script id="report-data" type="application/json">')[1].split('</script>')[0])
            administrative={a['id']:a for a in payload['tareas_administrativas']}
            self.assertEqual(administrative['6']['resultado'],'WARNING_ESTIMACION_ALTA')
            self.assertEqual(administrative['7']['resultado'],'SIN_ESTIMACION')
            self.assertEqual(administrative['8']['resultado'],'WARNING_ESTIMACION_ALTA')
            self.assertEqual(administrative['9']['resultado'],'REVISION_MANUAL')
            self.assertIn('adminrows',html)
            alerts=payload['disponibilidad_alertas']
            self.assertEqual(len(alerts),7)
            self.assertEqual({a['id'] for a in alerts},{'1','2','3','4'})
            self.assertTrue(all(a['horas'] is None for a in alerts))
            for reason in ('vacaciones','Permiso','Recurso inactivo','incapacidades'):
                self.assertTrue(any(reason in a['detalle'] for a in alerts),reason)
            self.assertIn('availabilityrows',html)
            book=load_workbook(html_paths[0].with_suffix('.xlsx'))
            self.assertEqual(book['Tareas administrativas'].max_row,5)
            self.assertEqual(book['Config tareas administrativas'].max_row,8)
            admin_warnings=[r for r in book['Alertas'].iter_rows(min_row=2,values_only=True) if r[0]=='ESTIMACION_ADMIN_ALTA']
            self.assertEqual({r[1] for r in admin_warnings},{'6','8'})
            self.assertEqual(book['Disponibilidad y tareas'].max_row,8)
            self.assertEqual({r[4] for r in book['Disponibilidad y tareas'].iter_rows(min_row=2,values_only=True)},{'1','2','3','4'})
            book.close()
