import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from test_reglas import report
from metricas import add_metrics
from dashboard import write_dashboard, generate_dashboard

class MetricasTest(unittest.TestCase):
    def test_only_identified_totals_and_all_types_matrix(self):
        book=Workbook(); book.remove(book.active)
        def sheet(name,headers,data):
            ws=book.create_sheet(name); ws.append(headers)
            for row in data: ws.append(row)
            return ws
        people=[{'documento':'1','nombre':'Ana','activo':True,'datos_base':{'Celula interna BD':'DB Tech'}}]
        def record(identifier,kind,hours,state='Active',person=0):
            return {'id':identifier,'persona':person,'sprint':'255','estado':state,'pendiente':state=='Active','horas':hours,'en_carga':kind!='Issue','asignado':'Ana','evc':'EVC','celula':'Celula'}
        records=[record('1','Task',8),record('2','HA',3,'Closed'),record('3','Issue',5,'Closed'),record('4','Task',99,person=None)]
        caps=[['1','Ana','255',True,1,0,8,0,0,68,8/68,'BAJO CAPACIDAD OBJETIVO',90,9,0,13,68,58.4375,48.875]]
        data=add_metrics(book,sheet,records,people,{'255':{'activo':True,'inicio':'2026-10-07','fin':'2026-10-20'}},caps,date(2026,10,8),[])
        self.assertEqual(book['Metricas sprint']['C2'].value,11)
        self.assertEqual(book['Metricas sprint']['D2'].value,3)
        self.assertEqual(book['Matriz asignadas']['F2'].value,11)
        self.assertEqual(book['Matriz total']['F2'].value,16)
        self.assertEqual(book['Carga recursos'].max_row,2)
        self.assertEqual(book['Historico recursos'].max_row,2)
        self.assertNotIn('Sprint',[c.value for c in book['Historico recursos'][1]])
        self.assertTrue(book['Equipos internos'].row_dimensions[3].hidden)
        self.assertEqual(book['Celulas y EVC']['E2'].value,11)
        self.assertEqual(book['Celulas y EVC']['D2'].value,3)
        self.assertTrue(book['Celulas y EVC'].row_dimensions[4].hidden)
        self.assertEqual(book['Detalle celulas EVC']['C2'].value,'255')
        self.assertEqual(data['indicadores'][4][1],3)
        self.assertEqual(len(data['recursos']),1)

    def test_dashboard_controls_exist(self):
        import re
        template=(Path(__file__).resolve().parents[1]/'src'/'dashboard.html').read_text(encoding='utf-8')
        controls=set(re.findall(r'id="([^" ]+)"',template))
        used=set(re.findall(r"\$\('([^']+)'\)",template))
        self.assertFalse(used-controls)

    def test_input_exactly_one_csv(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)
            with self.assertRaises(ValueError): report.select_csv(path)
            (path/'otro.CSV').write_text('ID\n1')
            (path/'ref.xlsx').write_text('ignored')
            self.assertEqual(report.select_csv(path).name,'otro.CSV')
            (path/'segundo.csv').write_text('ID\n2')
            with self.assertRaises(ValueError): report.select_csv(path)

    def test_holidays_two_dates_in_sprint_not_whole_month(self):
        person={'activo':True,'horas_dia':9,'dedicacion':1,'ausencias':[]}
        sprint={'activo':True,'sprint':'1','inicio':'2026-06-01','fin':'2026-06-15'}
        result=report.capacity_breakdown(person,sprint,['2026-06-08','2026-06-15','2026-06-29'],report.DEFAULT_TIMES)
        self.assertEqual(result['festivos'],18)
        self.assertEqual(result['disponibles'],68)

    def test_dashboard_serialization(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'dashboard.html'
            write_dashboard(path,{'nombre':'</script><script>alert(1)</script>'})
            html=path.read_text(encoding='utf-8')
            payload=html.split('<script id="report-data" type="application/json">')[1].split('</script>')[0]
            self.assertNotIn('<script>',payload)
            self.assertEqual(json.loads(payload)['nombre'],'</script><script>alert(1)</script>')

    def test_single_timestamped_dashboard_preserves_other_files(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            old = root / 'cargabilidad_anterior.html'
            write_dashboard(old, {'fecha': '2026-10-08'})
            unrelated = root / 'cargabilidad_notas.html'
            unrelated.write_text('<p>Notas</p>', encoding='utf-8')
            workbook = root / 'cargabilidad.xlsx'
            workbook.write_bytes(b'Excel existente')
            result = generate_dashboard(root, {'fecha': '2026-10-09'})
            self.assertRegex(result.name, r'^cargabilidad_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}\.html$')
            self.assertFalse(old.exists())
            self.assertEqual(unrelated.read_text(encoding='utf-8'), '<p>Notas</p>')
            self.assertEqual(workbook.read_bytes(), b'Excel existente')
            html = result.read_text(encoding='utf-8')
            self.assertIn('\u00faltimo sprint con horas registradas', html)
            self.assertNotIn('?ltimo', html)
            self.assertNotIn('selecci?n', html)
            self.assertIn('2026-10-09', html)
