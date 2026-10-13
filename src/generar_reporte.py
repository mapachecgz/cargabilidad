"""Generador local de cargabilidad. No conecta ni modifica Azure DevOps."""
import argparse
import csv
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TIMES = {
    'festivo_horas': 9,
    'actividades_por_sprint': {'daily': 6, 'planning': 2, 'retro': 1, 'reuniones_y_alineaciones': 4}
}


def normalize(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD', str(value or '')) if not unicodedata.combining(c)).casefold().split())

def identity(value):
    match = re.search(r'<([^>]+)>', value or '')
    return (value.split('<')[0].strip(), match.group(1).strip().lower() if match else '')

def read_csv(path):
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            with path.open(encoding=encoding, newline='') as f:
                sample = f.read(8192)
                f.seek(0)
                return list(csv.DictReader(f, delimiter=';' if sample.splitlines()[0].count(';') > sample.splitlines()[0].count(',') else ','))
        except UnicodeDecodeError:
            continue
    raise ValueError(f'No se pudo leer {path}')

def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')

def initialize(base, config):
    required = ['recursos.json', 'sprints.json', 'parametros.json', 'festivos.json', 'tiempos.json']
    missing = [name for name in required if not (config / name).is_file()]
    if missing:
        raise ValueError('Faltan configuraciones: ' + ', '.join(missing))

def select_csv(directory):
    files = sorted(p for p in directory.iterdir() if p.is_file() and p.suffix.lower() == '.csv')
    if len(files) != 1:
        raise ValueError(f'Entrada debe contener exactamente un CSV; encontrados: {len(files)}')
    return files[0]

def complete_sprints(sprints, rows):
    from datetime import timedelta
    known = [s for s in sprints if str(s['sprint']).isdigit()]
    missing = {m.group(1) for r in rows if (m := re.search(r'(\d+)\s*$', r['Iteration Level 3']))} - {str(s['sprint']) for s in sprints}
    if not missing: return []
    if not known: raise ValueError('No hay sprint numerico para inferir fechas')
    anchor = min(known, key=lambda s:int(s['sprint']))
    start = date.fromisoformat(anchor['inicio'])
    duration = (date.fromisoformat(anchor['fin'])-start).days+1
    for sprint in known:
        expected = start + timedelta(days=(int(sprint['sprint'])-int(anchor['sprint']))*duration)
        if date.fromisoformat(sprint['inicio']) != expected or date.fromisoformat(sprint['fin']) != expected+timedelta(days=duration-1):
            raise ValueError('Calendario irregular: no es posible inferir sprints sin revisar fechas')
    added=[]
    for identifier in sorted(missing,key=int):
        day=start+timedelta(days=(int(identifier)-int(anchor['sprint']))*duration)
        added.append({'sprint':identifier,'inicio':day.isoformat(),'fin':(day+timedelta(days=duration-1)).isoformat(),'activo':True,'origen':'inferido','patron_dias':duration})
    sprints.extend(added); sprints.sort(key=lambda s:int(s['sprint']))
    return added

def resource_absences(person):
    result = list(person.get('ausencias', []))
    for kind in ('vacaciones', 'incapacidades'):
        configured = person.get(kind, {})
        intervals = configured if isinstance(configured, list) else [configured]
        for interval in intervals:
            if not interval or not interval.get('inicio') and not interval.get('fin'): continue
            if not interval.get('inicio') or not interval.get('fin'):
                raise ValueError(f'Faltan fechas de {kind} en {person["nombre"]}')
            result.append({**interval, 'motivo':kind})
    for interval in result:
        if date.fromisoformat(interval['fin']) < date.fromisoformat(interval['inicio']):
            raise ValueError('Ausencia con fechas invertidas')
    return result

def unavailable(person, day):
    return not person['activo'] or any(date.fromisoformat(a['inicio']) <= day <= date.fromisoformat(a['fin']) for a in resource_absences(person))

def availability_reasons(person, day):
    reasons = []
    if not person['activo']:
        reasons.append('Recurso inactivo')
    for absence in resource_absences(person):
        if absence['inicio'] <= day.isoformat() <= absence['fin']:
            reasons.append(f"{absence.get('motivo') or 'Ausencia'}: {absence['inicio']} a {absence['fin']}")
    return reasons


def net_hours(gross, holidays, absences, activities):
    return max(0, gross - holidays - absences - activities)

def capacity_breakdown(person, sprint, holidays, times=None):
    from datetime import timedelta
    start, end = date.fromisoformat(sprint['inicio']), date.fromisoformat(sprint['fin'])
    if end < start:
        raise ValueError('Sprint con fechas invertidas: ' + sprint['sprint'])
    if not person['activo'] or not sprint['activo']:
        return {'brutas': 0, 'festivos': 0, 'ausencias': 0, 'actividades': 0, 'disponibles': 0}
    days = [start + timedelta(days=i) for i in range((end-start).days+1) if (start+timedelta(days=i)).weekday() < 5]
    daily = person['horas_dia'] * person['dedicacion']
    gross = len(days) * daily
    holiday_days = {d for d in days if d.isoformat() in holidays}
    holiday_hours = len(holiday_days) * (times['festivo_horas'] * person['dedicacion'] if times else daily)
    absent_hours = sum(daily for d in days if d not in holiday_days and unavailable(person, d))
    activities = sum(times['actividades_por_sprint'].values()) if times else 0
    return {'brutas': gross, 'festivos': holiday_hours, 'ausencias': absent_hours, 'actividades': activities,
            'disponibles': net_hours(gross, holiday_hours, absent_hours, activities)}

def capacity(person, sprint, holidays, times=None):
    return capacity_breakdown(person, sprint, holidays, times)['disponibles']

def thresholds(available, times):
    # Proporciones acordadas: todos los limites varian con la disponibilidad neta.
    return {'exceso_desde': available, 'punto_medio': available * 55 / 64, 'capacidad': available * 46 / 64}

def load_status(hours, available, missing, limits):
    if available is None: return 'SIN FECHAS'
    if available == 0: return 'SIN CAPACIDAD'
    if hours >= limits['exceso_desde']: return 'EXCESO'
    if missing: return 'ESTIMACION INCOMPLETA'
    if hours >= limits['punto_medio']: return 'SOBRE PUNTO MEDIO'
    if hours >= limits['capacidad']: return 'SOBRE CAPACIDAD OBJETIVO'
    return 'BAJO CAPACIDAD OBJETIVO'


def number(value):
    return float(str(value).replace(',', '.')) if value not in ('', None) else None

def previous_pending(records, people, sprints, cutoff):
    result=[]
    for r in records:
        sprint=sprints.get(r['sprint'])
        if r['persona'] is None or not r['pendiente'] or not sprint or sprint['fin'] >= cutoff.isoformat(): continue
        person=people[r['persona']]
        result.append({'id':r['id'],'titulo':r.get('titulo',''),'tipo':r['tipo'],'estado':r['estado'],'sprint':r['sprint'],'fin':sprint['fin'],'dias':(cutoff-date.fromisoformat(sprint['fin'])).days,'horas':r['horas'],'nombre':person['nombre'],'activo':person['activo'],'celula':person['datos_base'].get('Celula interna BD') or 'Sin celula','evc':r['evc'],'celula_azure':r['celula']})
    return sorted(result,key=lambda r:(r['fin'],r['nombre'],r['id']))

def run(base, config, output, cutoff):
    initialize(base, config)
    people = json.loads((config / 'recursos.json').read_text(encoding='utf-8'))
    sprints = json.loads((config / 'sprints.json').read_text(encoding='utf-8'))
    params = json.loads((config / 'parametros.json').read_text(encoding='utf-8'))
    holidays = json.loads((config / 'festivos.json').read_text(encoding='utf-8'))['fechas']
    holidays = sorted({date.fromisoformat(day).isoformat() for day in holidays})
    times = json.loads((config / 'tiempos.json').read_text(encoding='utf-8'))
    from administrativos import load_rules, classify
    administrative_settings = load_rules(config / 'tareas_administrativas.json')
    values = [times['festivo_horas'], *times['actividades_por_sprint'].values()]
    if any(not isinstance(v, (int, float)) or v < 0 for v in values):
        raise ValueError('Los tiempos deben ser numeros no negativos')
    source = select_csv(base)
    rows = read_csv(source)
    required = {'ID', 'Parent', 'Title', 'Work Item Type', 'Assigned To', 'State', 'Area Level 4', 'Area Level 3', 'Iteration Level 3', 'Original Estimate', 'Story Points'}
    if not rows or not required.issubset(rows[0]):
        raise ValueError('CSV vacio o faltan columnas obligatorias')
    inferred = complete_sprints(sprints, rows)
    if inferred: save_json(config / 'sprints.json', sprints)
    index = {}
    for row in rows:
        if row['ID'] in index:
            raise ValueError('ID duplicado: ' + row['ID'])
        index[row['ID']] = row
    matches = defaultdict(set)
    for i, p in enumerate(people):
        for key in [p['nombre'], p.get('correo', '')] + p.get('alias', []):
            if key:
                matches[normalize(key)].add(i)
    sprint_map = {s['sprint']: s for s in sprints}
    if len(sprint_map) != len(sprints):
        raise ValueError('Sprint duplicado en configuracion')
    alerts, detail = [], []
    records = []
    administrative_rows = []
    load = defaultdict(lambda: [0, 0, 0.0, 0, 0.0])
    groups = defaultdict(lambda: [0, 0, 0.0, 0.0])
    def alert(row, kind, message):
        alerts.append([kind, row['ID'], row['Parent'], row['Work Item Type'], row['Assigned To'], row['State'], row['Iteration Level 3'], row['Area Level 4'], row['Area Level 3'], message])
    for row in rows:
        name, email = identity(row['Assigned To'])
        candidates = matches.get(normalize(email), set()) if email else set()
        if not candidates:
            candidates = matches.get(normalize(name), set())
        person_id = next(iter(candidates)) if len(candidates) == 1 else None
        person = people[person_id] if person_id is not None else None
        pending = row['State'] not in params['estados_finales']
        m = re.search(r'(\d+)\s*$', row['Iteration Level 3'])
        sprint_id = m.group(1) if m else ''
        sprint = sprint_map.get(sprint_id)
        ancestors, seen, parent = [], {row['ID']}, row['Parent']
        while parent:
            if parent in seen:
                alert(row, 'CICLO_PARENT', 'La jerarquia contiene un ciclo')
                break
            seen.add(parent)
            if parent not in index:
                alert(row, 'PARENT_NO_EXPORTADO', 'El padre no esta incluido en el CSV')
                break
            ancestors.append(parent)
            parent = index[parent]['Parent']
        if not sprint:
            alert(row, 'SPRINT_SIN_CONFIGURACION', 'Sin fechas conocidas; no se inventa capacidad')
        if pending:
            if not name:
                alert(row, 'SIN_ASIGNAR', 'Elemento pendiente sin responsable')
            elif person is None:
                alert(row, 'RECURSO_AMBIGUO' if len(candidates)>1 else 'RECURSO_NO_CATALOGADO', 'Revise correo o alias en recursos.json')
            if person and unavailable(person, cutoff):
                alert(row, 'RECURSO_NO_DISPONIBLE', 'Fecha de corte ' + cutoff.isoformat() + ': ' + '; '.join(availability_reasons(person, cutoff)) + '. Tiene un elemento abierto; revisar cierre o reasignacion.')
            if sprint and date.fromisoformat(sprint['fin']) < cutoff:
                alert(row, 'PENDIENTE_SPRINT_FINALIZADO', 'Elemento pendiente en un sprint finalizado')
            if sprint and not sprint['activo']:
                alert(row, 'SPRINT_INACTIVO', 'Elemento pendiente en sprint desactivado')
            if person and sprint and person['activo'] and any(a['inicio'] <= sprint['fin'] and a['fin'] >= sprint['inicio'] for a in resource_absences(person)):
                overlaps = [a for a in resource_absences(person) if a['inicio'] <= sprint['fin'] and a['fin'] >= sprint['inicio']]
                alert(row, 'AUSENCIA_EN_SPRINT', '; '.join(f"{a.get('motivo') or 'Ausencia'}: {a['inicio']} a {a['fin']}" for a in overlaps) + '. Coincide con el sprint; revisar asignacion. No se conocen fechas de ejecucion por tarea.')
        estimate, points = None, None
        for column in ('Original Estimate', 'Story Points'):
            try:
                parsed = number(row[column])
                if parsed is not None and parsed < 0:
                    raise ValueError()
                if column == 'Original Estimate': estimate = parsed
                else: points = parsed
            except ValueError:
                alert(row, 'VALOR_INVALIDO', column + ': ' + row[column])
        evaluation = classify(row['Title'], row['Work Item Type'], row['State'], points, administrative_settings, params['estados_finales'])
        if evaluation:
            administrative_rows.append({**evaluation, 'id':row['ID'], 'sprint':sprint_id, 'tipo':row['Work Item Type'],
                'titulo':row['Title'], 'estado':row['State'], 'horas':points,
                'nombre':person['nombre'] if person else name or 'Sin asignar', 'activo':person['activo'] if person else None,
                'celula':(person['datos_base'].get('Celula interna BD') or 'Sin celula') if person else 'No identificado',
                'identificado':person is not None})
            if evaluation['resultado'] == 'WARNING_ESTIMACION_ALTA':
                alert(row, 'ESTIMACION_ADMIN_ALTA', evaluation['categoria'] + ': ' + evaluation['detalle'])
        is_task = row['Work Item Type'] in params['tipos_carga']
        is_ha = row['Work Item Type'] in params['tipos_ha']
        records.append({'persona':person_id,'sprint':sprint_id,'estado':row['State'],'pendiente':pending,'horas':points,'en_carga':is_task or is_ha,'asignado':row['Assigned To'],'id':row['ID'],'tipo':row['Work Item Type'],'titulo':row['Title'],'celula':row['Area Level 4'],'evc':row['Area Level 3']})
        if pending and (is_task or is_ha) and points is None:
            alert(row, 'SIN_ESTIMACION', 'Carga incompleta: elemento pendiente sin Story Points (horas) valido')
        if person is not None:
            bucket = load[(person_id, sprint_id)]
            bucket[0] += int(pending and is_task)
            bucket[1] += int(pending and is_ha)
            bucket[2] += (points or 0) if pending and (is_task or is_ha) else 0
            bucket[3] += int(pending and (is_task or is_ha) and points is None)
            bucket[4] += (points or 0) if pending and is_ha else 0
        bucket = groups[(row['Area Level 3'], row['Area Level 4'], sprint_id)]
        bucket[0] += 1
        bucket[1] += int(pending)
        bucket[2] += (points or 0) if pending and is_task else 0
        bucket[3] += (points or 0) if pending and is_ha else 0
        detail.append(list(row.values()) + [row['Area Level 4'], row['Area Level 3'], name, email, person['documento'] if person else '', person['nombre'] if person else '', 'Si' if pending else 'No', sprint_id, ' > '.join(reversed(ancestors)), estimate, points])
    resource_rows = []
    for i, person in enumerate(people):
        ids = set(sprint_map) | {s for p,s in load if p == i}
        for sid in sorted(ids, key=lambda s: int(s) if s.isdigit() else -1):
            tasks, ha, hours, missing, ha_hours = load[(i, sid)]
            breakdown = capacity_breakdown(person, sprint_map[sid], holidays, times) if sid in sprint_map else None
            cap = breakdown['disponibles'] if breakdown else None
            limits = thresholds(cap, times) if cap is not None else dict.fromkeys(['exceso_desde','punto_medio','capacidad'])
            usage = hours/cap if cap else None
            status = load_status(hours, cap, missing, limits)
            if cap is not None and hours > 0 and (cap == 0 or hours >= limits['exceso_desde']):
                alerts.append(['SOBRECARGA', '', '', '', person['nombre'], '', sid, '', '', f'{hours} horas pendientes / {cap} horas disponibles; exceso desde {limits["exceso_desde"]}'])
            resource_rows.append([person['documento'],person['nombre'],sid,person['activo'],tasks,ha,hours,missing,ha_hours,cap,usage,status] +
                                 [breakdown[k] if breakdown else None for k in ('brutas','festivos','ausencias','actividades')] +
                                 [limits[k] for k in ('exceso_desde','punto_medio','capacidad')])

    book = Workbook()
    book.remove(book.active)
    def sheet(title, headers, data):
        ws = book.create_sheet(title)
        ws.append(headers)
        for values in data:
            ws.append(values)
        for cells in ws:
            for cell in cells:
                if cell.data_type == 'f':
                    cell.data_type = 's'
        ws.freeze_panes = 'A2'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = 'C8102E'
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='8F0B17')
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[1].height = 32
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(55, max(14, max(len(str(c.value or '')) for c in list(col)[:120])+2))
        if data:
            table_id = len(book.worksheets)
            existing_tables = {name for existing in book for name in existing.tables}
            while 'Tabla'+str(table_id) in existing_tables: table_id += 1
            table = Table(displayName='Tabla'+str(table_id),ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(table)
        else:
            ws.auto_filter.ref = ws.dimensions
        return ws
    sheet('Detalle',list(rows[0])+['Celula Azure','EVC Azure','Nombre asignado','Correo asignado','Documento talento','Talento vinculado','Pendiente','Sprint ID','Ancestros exportados','Horas originales numericas','Horas Story Points numericas'],detail)
    sheet('Alertas',['Tipo','ID','Parent','Tipo elemento','Responsable','Estado','Sprint','Celula Azure','EVC Azure','Detalle'],alerts)
    headers = list(dict.fromkeys(k for p in people for k in p['datos_base']))
    sheet('Base talentos',headers,[[p['datos_base'].get(k) for k in headers] for p in people])
    resources_ws = sheet('Config recursos',['Documento','Talento','Correo','Alias','Activo','Horas dia','Dedicacion al equipo','Ausencias','Vacaciones','Incapacidades'],[[p['documento'],p['nombre'],p.get('correo',''),', '.join(p.get('alias',[])),p['activo'],p['horas_dia'],p['dedicacion'],json.dumps(p.get('ausencias',[]),ensure_ascii=False),json.dumps(p.get('vacaciones',{}),ensure_ascii=False),json.dumps(p.get('incapacidades',{}),ensure_ascii=False)] for p in people])
    for cell in resources_ws['G'][1:]: cell.number_format = '0%'
    sheet('Config sprints',['Sprint','Inicio','Fin','Activo','Origen'],[[s['sprint'],s['inicio'],s['fin'],s['activo'],s.get('origen','configurado')] for s in sprints])
    sheet('Config festivos',['Fecha'],[[day] for day in holidays])
    sheet('Config tiempos',['Grupo','Concepto','Horas'],[['Descuentos','Por festivo',times['festivo_horas']]] + [['Actividades por sprint',k,v] for k,v in times['actividades_por_sprint'].items()])
    sheet('Tareas administrativas', ['Recurso','Identificado','Activo','Celula interna','Sprint','ID','Tipo elemento','Titulo','Estado','Horas registradas','Reglas','Categoria','Horas referencia','Umbral warning','Diferencia referencia','Veces referencia','Resultado','Detalle'],
        [[a[k] for k in ('nombre','identificado','activo','celula','sprint','id','tipo','titulo','estado','horas','reglas','categoria','referencia','umbral','diferencia','multiplicador','resultado','detalle')] for a in administrative_rows])
    sheet('Config tareas administrativas', ['ID regla','Categoria','Activa','Patrones','Exclusiones','Horas referencia','Umbral warning'],
        [[r['id'],r['nombre'],r['activo'],json.dumps(r['patrones'],ensure_ascii=False),json.dumps(r.get('excluir_patrones',[]),ensure_ascii=False),r['horas_referencia'],r['horas_maximas_sin_warning']] for r in administrative_settings['reglas']])
    from metricas import add_metrics
    payload = add_metrics(book, sheet, records, people, sprint_map, resource_rows, cutoff, alerts)
    payload['tareas_administrativas'] = administrative_rows
    payload['administrativas_activas'] = administrative_settings['activo']
    from vistas import reorganize
    reorganize(book, payload, records, people)
    linked_ids = {r['id'] for r in records if r['persona'] is not None}
    included_alerts = {'RECURSO_NO_DISPONIBLE','AUSENCIA_EN_SPRINT','SOBRECARGA','SIN_ESTIMACION','PENDIENTE_SPRINT_FINALIZADO'}
    payload['alertas'] = [{'tipo':a[0],'id':a[1],'recurso':identity(a[4])[0],'sprint':re.search(r'(\d+)\s*$',a[6]).group(1) if re.search(r'(\d+)\s*$',a[6]) else '', 'detalle':a[9]} for a in alerts if a[0] in included_alerts and (a[1] in linked_ids or a[0]=='SOBRECARGA')]
    records_by_id = {r['id']: r for r in records}
    availability_alerts = []
    for a in alerts:
        if a[0] not in {'RECURSO_NO_DISPONIBLE', 'AUSENCIA_EN_SPRINT'}:
            continue
        record = records_by_id[a[1]]
        person = people[record['persona']]
        availability_alerts.append({'nombre':person['nombre'], 'activo':person['activo'],
            'celula':person['datos_base'].get('Celula interna BD') or 'Sin celula',
            'sprint':record['sprint'], 'id':record['id'], 'tipo':record['tipo'],
            'titulo':record['titulo'], 'estado':record['estado'], 'horas':record['horas'],
            'alerta':a[0], 'detalle':a[9]})
    payload['disponibilidad_alertas'] = availability_alerts
    sheet('Disponibilidad y tareas', ['Recurso','Activo','Celula interna','Sprint','ID','Tipo elemento','Titulo','Estado','Horas registradas','Tipo alerta','Motivo y alcance'],
        [[a[k] for k in ('nombre','activo','celula','sprint','id','tipo','titulo','estado','horas','alerta','detalle')] for a in availability_alerts])
    payload['pendientes_anteriores'] = previous_pending(records, people, sprint_map, cutoff)
    finished = [s for s in sprints if s['fin'] < cutoff.isoformat()]
    payload['ultimo_finalizado'] = max(finished,key=lambda s:s['fin'])['sprint'] if finished else None
    backlog=payload['pendientes_anteriores']
    sheet('Pendientes anteriores',['Recurso','Sprint','Fin sprint','Dias desde fin','ID','Tipo','Titulo','Estado','Horas registradas','Activo','Celula interna','Celula Azure','EVC Azure'],[[r[k] for k in ('nombre','sprint','fin','dias','id','tipo','titulo','estado','horas','activo','celula','celula_azure','evc')] for r in backlog])
    timestamp = datetime.now(timezone(timedelta(hours=-5))).strftime('%Y-%m-%d_%H-%M-%S')
    output = output.parent / f'cargabilidad_{timestamp}.xlsx'
    output.parent.mkdir(parents=True,exist_ok=True)
    book.save(output)
    from dashboard import generate_dashboard
    dashboard_path = generate_dashboard(output.parent, payload, timestamp=timestamp)
    print(json.dumps({'archivo':str(output),'entrada':str(source),'dashboard':str(dashboard_path),'filas_csv':len(rows),'talentos':len(people),'alertas':len(alerts),'tipos_alerta':dict(Counter(a[0] for a in alerts))},ensure_ascii=False,indent=2))
    return output

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--entrada','--base',dest='base',type=Path,default=ROOT/'entrada',help='Directorio con exactamente un CSV')
    parser.add_argument('--config',type=Path,default=ROOT/'config')
    output_options = parser.add_mutually_exclusive_group()
    output_options.add_argument('--directorio-salida',type=Path,help='Directorio para el Excel y HTML con fecha y hora')
    output_options.add_argument('--salida',type=Path,default=ROOT/'salida'/'cargabilidad.xlsx',help='Ruta cuyo directorio se usa para los reportes cargabilidad con fecha y hora')
    parser.add_argument('--fecha-corte',type=date.fromisoformat,default=datetime.now(timezone(timedelta(hours=-5))).date())
    args = parser.parse_args()
    run(args.base,args.config,(args.directorio_salida/'cargabilidad.xlsx') if args.directorio_salida else args.salida,args.fecha_corte)
