"""Vistas Excel con subtotales y detalle desplegable."""
from collections import defaultdict
from copy import copy
from openpyxl.styles import Font, PatternFill


def reorganize(book, payload, records, people):
    def rebuild(name, headers, blocks):
        ws=book[name]
        ws._tables.clear(); ws.auto_filter.ref=None
        ws.delete_rows(1,ws.max_row)
        ws.row_dimensions.clear()
        ws.append(headers)
        ws.row_dimensions[1].height=32; ws.freeze_panes='A2'
        ws.sheet_properties.outlinePr.summaryBelow=False
        for c in ws[1]:
            c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='8F0B17')
        for summary,children in blocks:
            top=ws.max_row+1; ws.append(summary)
            for c in ws[top]: c.font=Font(bold=True,color='8F0B17'); c.fill=PatternFill('solid',fgColor='F2F2F0')
            first=ws.max_row+1
            for child in children: ws.append(child)
            if children:
                ws.row_dimensions.group(first,ws.max_row,outline_level=1,hidden=True)
                ws.row_dimensions[top].collapsed=True
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(48,max(16,max(len(str(c.value or '')) for c in list(col)[:120])+2))
        return ws
    data=payload['recursos']; names={p['nombre']:p for p in people}
    for name in ('Carga recursos','Historico recursos'):
        original=list(book[name].iter_rows(values_only=True)); headers=list(original[0])+['Sprint']
        blocks=[]
        for summary in original[1:]:
            resource=summary[1]; person=names[resource]; children=[]
            for d in data:
                if d['nombre']!=resource or d['sprint'] not in payload['sprintsCSV']: continue
                if name=='Carga recursos':
                    child=[None,'Sprint '+(d['sprint'] or 'Sin identificar'),person['activo'],d['horas'],d['cerradas'],d['pendientes'],d['otros'],d['disponibles'],d['pendientes']/d['disponibles'] if d['disponibles'] else None,d['faltantes'],int(d['disponibles'] is None),d['sprint'] or 'Sin sprint']
                else:
                    missing=sum(r['horas'] is None for r in records if r['persona'] is not None and people[r['persona']]['nombre']==resource and r['sprint']==d['sprint'] and r['en_carga'])
                    child=[None,'Sprint '+(d['sprint'] or 'Sin identificar'),d['horas'],d['cerradas'],d['pendientes'],d['otros'],missing,d['total'],d['sprint'] or 'Sin sprint']
                children.append(child)
            blocks.append((list(summary)+['TOTAL RECURSO'],children))
        ws=rebuild(name,headers,blocks)
        if name=='Carga recursos':
            for c in ws['I'][1:]: c.number_format='0.0%'
    ws=book['Metricas sprint']; original=list(ws.iter_rows(values_only=True)); blocks=[]
    for row in original[1:]:
        sid=row[0]; children=[]
        for d in data:
            if d['sprint']!=sid: continue
            related=[r for r in records if r['persona'] is not None and people[r['persona']]['nombre']==d['nombre'] and r['sprint']==sid and r['en_carga']]
            missing=sum(r['horas'] is None for r in related)
            c=d['disponibles']; excess=int(d['activo'] and d['pendientes']>0 and c is not None and d['pendientes']>=c)
            child=[d['nombre'],len(related),d['horas'],d['cerradas'],d['pendientes'],d['otros'],missing,c,d['pendientes']/c if c else None,d['horas']/c if c else None,excess,int(d['activo'] and bool(c) and d['horas']==0),int(not d['activo'] and d['pendientes']>0),d['cerradas']/d['horas'] if d['horas'] else None,d['faltantes'],'Con registros' if related else 'Sin registros',d['nombre']]
            children.append(child)
        blocks.append((list(row)+['TOTAL EQUIPO'],children))
    ws=rebuild('Metricas sprint',list(original[0])+['Recurso'],blocks)
    for col in ('I','J','N'):
        for c in ws[col][1:]: c.number_format='0.0%'
    # Distinct EVC/cell/sprint groups, resources are children, not duplicate totals.
    groups=defaultdict(lambda:defaultdict(lambda:[0,0.,0.,0.,0.]))
    for r in records:
        if r['persona'] is None: continue
        v=groups[(r['evc'],r['celula'],r['sprint'])][people[r['persona']]['nombre']]
        v[0]+=1; v[4]+=r['horas'] or 0
        if r['en_carga']:
            v[1]+=r['horas'] or 0
            v[2]+=(r['horas'] or 0) if r['estado']=='Closed' else 0
            v[3]+=(r['horas'] or 0) if r['pendiente'] else 0
    blocks=[]
    for k,members in sorted(groups.items()):
        totals=[sum(v[j] for v in members.values()) for j in range(5)]
        blocks.append((list(k)+totals+['TOTAL GRUPO'],[[n,'','']+v+[n] for n,v in sorted(members.items())]))
    rebuild('Detalle celulas EVC',['EVC Azure','Celula Azure','Sprint','Registros Azure','Horas asignadas','Horas cerradas','Horas pendientes','Horas todos los tipos','Recurso'],blocks)
    original=list(book['Equipos internos'].iter_rows(values_only=True)); blocks=[]
    for row in original[1:]:
        if row[1]!='TOTAL CELULA': continue
        cell=row[0]; children=[]
        for d in data:
            if d['celula']==cell and d['sprint'] in payload['sprintsCSV']:
                children.append([cell,d['nombre'],int(d['activo']),d['horas'],d['cerradas'],d['pendientes'],d['disponibles'],d['sprint'] or 'Sin sprint'])
        blocks.append((list(row)+['TODOS'],children))
    rebuild('Equipos internos',list(original[0])+['Sprint'],blocks)
    # Every additive dashboard indicator exposes its contributors.
    original=list(book['Tablero'].iter_rows(values_only=True)); blocks=[]
    for indicator,value in original[1:]:
        contributors=[]
        if indicator in ('Recursos configurados','Recursos activos','Recursos inactivos'):
            contributors=[[p['nombre'],1,'Configuracion'] for p in people if indicator=='Recursos configurados' or p['activo']==(indicator=='Recursos activos')]
        elif indicator=='Elementos de recursos identificados':
            for i,p in enumerate(people): contributors.append([p['nombre'],sum(r['persona']==i for r in records),'Todo el CSV'])
        elif indicator in ('Horas asignadas (tipos de carga)','Horas cerradas','Horas pendientes'):
            key={'Horas asignadas (tipos de carga)':'horas','Horas cerradas':'cerradas','Horas pendientes':'pendientes'}[indicator]
            contributors=[[p['nombre'],sum(d[key] for d in data if d['nombre']==p['nombre'] and d['sprint'] in payload['sprintsCSV']),'Todos los sprints del CSV'] for p in people]
        elif indicator.startswith('Sprint ') and isinstance(value,(int,float)):
            sid=indicator.split(':')[0].split()[1]; selection=[d for d in data if d['sprint']==sid]
            if 'horas registradas' in indicator: key='horas'
            elif 'horas de trabajo pendientes' in indicator: key='pendientes'
            elif 'horas disponibles' in indicator: key='disponibles'
            elif 'sin estimacion' in indicator: key='faltantes'
            else: key=None
            if key: contributors=[[d['nombre'],d[key],sid] for d in selection]
            elif 'exceso' in indicator: contributors=[[d['nombre'],1,sid] for d in selection if d['activo'] and d['pendientes']>0 and d['disponibles'] is not None and d['pendientes']>=d['disponibles']]
            elif 'activos sin' in indicator: contributors=[[d['nombre'],1,sid] for d in selection if d['activo'] and d['disponibles'] and d['horas']==0]
            elif 'inactivos' in indicator: contributors=[[d['nombre'],1,sid] for d in selection if not d['activo'] and d['pendientes']>0]
        blocks.append(([indicator,value,'TOTAL'],contributors))
    rebuild('Tablero',['Indicador / Recurso','Valor','Periodo'],blocks)
    # Tablero: recurso -> sprint -> tipo de elemento.
    ws=book['Tablero']
    ws.cell(1,4,'Sprint'); ws.cell(1,4).font=Font(bold=True,color='FFFFFF'); ws.cell(1,4).fill=PatternFill('solid',fgColor='8F0B17')
    header=next((r for r in range(2,ws.max_row+1) if ws.cell(r,1).value=='Elementos de recursos identificados'),None)
    if header:
        insert=header+1
        for i,p in enumerate(people):
            per_sprint=defaultdict(lambda:defaultdict(int))
            for rec in records:
                if rec['persona']==i:
                    per_sprint[rec['sprint']][rec.get('tipo','Sin tipo')]+=1
            insert+=1
            for sprint,types in sorted(per_sprint.items(),key=lambda item:(0,int(item[0])) if item[0].isdigit() else (1,item[0])):
                ws.insert_rows(insert)
                ws.cell(insert,1,'Sprint '+(sprint or 'Sin identificar')); ws.cell(insert,2,sum(types.values())); ws.cell(insert,3,'Subtotal sprint'); ws.cell(insert,4,sprint or 'Sin sprint')
                insert+=1
                for kind,count in sorted(types.items()):
                    ws.insert_rows(insert)
                    ws.cell(insert,1,kind); ws.cell(insert,2,count); ws.cell(insert,3,'Tipo de elemento'); ws.cell(insert,4,sprint or 'Sin sprint')
                    insert+=1
        ws.row_dimensions.clear()
        current_header=None; resource_header=None; sprint_header=None
        for row in range(2,ws.max_row+1):
            period=ws.cell(row,3).value
            if period=='TOTAL':
                current_header=row; resource_header=None; sprint_header=None
            elif current_header:
                level=3 if period=='Tipo de elemento' else 2 if period=='Subtotal sprint' else 1
                ws.row_dimensions[row].hidden=True; ws.row_dimensions[row].outlineLevel=level
                ws.row_dimensions[current_header].collapsed=True
                if level==1: resource_header=row; sprint_header=None
                elif level==2:
                    sprint_header=row
                    if resource_header: ws.row_dimensions[resource_header].collapsed=True
                elif sprint_header: ws.row_dimensions[sprint_header].collapsed=True
    # Los indicadores historicos de horas tambien exponen recurso -> sprint.
    entries=[]; active_indicator=None
    hour_keys={'Horas asignadas (tipos de carga)':'horas','Horas cerradas':'cerradas','Horas pendientes':'pendientes'}
    for row in range(2,ws.max_row+1):
        values=[ws.cell(row,col).value for col in range(1,5)]
        if values[2]=='TOTAL': active_indicator=values[0]
        if active_indicator and active_indicator.startswith('Sprint '):
            values[3]=active_indicator.split(':')[0].split()[1]
        level=ws.row_dimensions[row].outlineLevel
        collapsed=ws.row_dimensions[row].collapsed
        entries.append((values,level,collapsed))
        if active_indicator in hour_keys and values[2]=='Todos los sprints del CSV':
            key=hour_keys[active_indicator]
            selection=[d for d in data if d['nombre']==values[0] and d['sprint'] in payload['sprintsCSV']]
            # Conservar los sprints con elementos para ese recurso, incluso si sus horas son cero.
            for d in selection:
                if not any(rec['persona'] is not None and people[rec['persona']]['nombre']==values[0] and rec['sprint']==d['sprint'] for rec in records): continue
                entries.append((['Sprint '+(d['sprint'] or 'Sin identificar'),d[key],'Detalle horas por sprint',d['sprint'] or 'Sin sprint'],2,False))
            if len(entries) and selection:
                resource_index=next(i for i in range(len(entries)-1,-1,-1) if entries[i][0] is values)
                entries[resource_index]=(values,level,True)
    ws.delete_rows(2,ws.max_row); ws.row_dimensions.clear()
    for index,(values,level,collapsed) in enumerate(entries,2):
        ws.append(values); ws.row_dimensions[index].outlineLevel=level
        ws.row_dimensions[index].hidden=bool(level); ws.row_dimensions[index].collapsed=collapsed
        if level==0:
            for cell in ws[index]: cell.font=Font(bold=True,color='8F0B17'); cell.fill=PatternFill('solid',fgColor='F2F2F0')
    ws.column_dimensions['D'].width=18
    # Equipos internos: celula -> recurso -> sprint, labels appear in column A.
    cells=sorted({p['datos_base'].get('Celula interna BD') or 'Sin celula' for p in people})
    rows=[]; levels=[]
    for cell in cells:
        subset=[d for d in data if d['celula']==cell and d['sprint'] in payload['sprintsCSV']]
        totals=[sum(d[k] or 0 for d in subset) for k in ('horas','cerradas','pendientes','disponibles')]
        rows.append([cell,'TOTAL CELULA',len({d['nombre'] for d in subset}),*totals,'TODOS']); levels.append(0)
        for name in sorted({d['nombre'] for d in subset}):
            own=[d for d in subset if d['nombre']==name]
            sums=[sum(d[k] or 0 for d in own) for k in ('horas','cerradas','pendientes','disponibles')]
            rows.append([name,name,int(names[name]['activo']),*sums,'TOTAL RECURSO']); levels.append(1)
            for d in own:
                rows.append(['Sprint '+(d['sprint'] or 'Sin identificar'),name,int(d['activo']),d['horas'],d['cerradas'],d['pendientes'],d['disponibles'],d['sprint']]); levels.append(2)
    def flat_outline(name,headers,rows,levels):
        ws=book[name]; ws._tables.clear(); ws.auto_filter.ref=None
        ws.delete_rows(1,ws.max_row); ws.row_dimensions.clear(); ws.append(headers)
        ws.sheet_properties.outlinePr.summaryBelow=False
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='8F0B17')
        for idx,(row,level) in enumerate(zip(rows,levels),2):
            ws.append(row); ws.row_dimensions[idx].outlineLevel=level; ws.row_dimensions[idx].hidden=bool(level)
            if idx-2+1<len(levels) and levels[idx-2+1]>level: ws.row_dimensions[idx].collapsed=True
        ws.column_dimensions['A'].width=52
        return ws
    flat_outline('Equipos internos',['Celula / recurso / sprint','Recurso','Activo / cantidad','Horas asignadas','Horas cerradas','Horas pendientes','Disponibles','Sprint'],rows,levels)
    # Cell/EVC hierarchy: existing subtotals and sprint leaves -> resource leaves.
    ws=book['Celulas y EVC']; existing=list(ws.iter_rows(values_only=True)); rows=[]; levels=[]
    for row in existing[1:]:
        evc,cell,sid=row[:3]
        if cell=='TOTAL EVC': rows.append(list(row)+['']); levels.append(0)
        elif sid=='TOTAL CELULA': rows.append([cell,*row[1:],'']); levels.append(1)
        else:
            rows.append(['Sprint '+str(sid),*row[1:],'']); levels.append(2)
            raw_sid='' if sid=='Sin sprint' else str(sid)
            for n,v in sorted(groups[(evc,cell,raw_sid)].items()):
                rows.append([n,cell,sid,*v,n]); levels.append(3)
    flat_outline('Celulas y EVC',['EVC / celula / sprint / recurso',*existing[0][1:],'Recurso'],rows,levels)
    # Detalle de organizacion: el sprint forma parte del esquema, no una columna.
    detail_rows=[]; detail_levels=[]
    def summed(keys):
        return [sum(v[j] for key in keys for v in groups[key].values()) for j in range(5)]
    for evc in sorted({key[0] for key in groups}):
        evc_keys=[key for key in groups if key[0]==evc]
        detail_rows.append([evc,*summed(evc_keys),'EVC']); detail_levels.append(0)
        for cell in sorted({key[1] for key in evc_keys}):
            cell_keys=[key for key in evc_keys if key[1]==cell]
            detail_rows.append([cell,*summed(cell_keys),'Celula']); detail_levels.append(1)
            for key in sorted(cell_keys,key=lambda key:(0,int(key[2])) if key[2].isdigit() else (1,key[2])):
                detail_rows.append(['Sprint '+(key[2] or 'Sin identificar'),*summed([key]),'Sprint']); detail_levels.append(2)
                for resource,values in sorted(groups[key].items()):
                    detail_rows.append([resource,*values,'Recurso']); detail_levels.append(3)
    flat_outline('Detalle celulas EVC',['EVC / celula / sprint / recurso','Registros Azure','Horas asignadas','Horas cerradas','Horas pendientes','Horas todos los tipos','Nivel'],detail_rows,detail_levels)
    # A single matrix: scope/state groups with resources as children.
    source=book['Matriz total']; headers=[c.value for c in source[1]]
    rows=[]; levels=[]
    states=sorted({r['estado'] for r in records if r['persona'] is not None})
    categories=[('TOTAL - todos los tipos',None,False)]+[(state,state,False) for state in states]
    for label,state,only_load in categories:
        detail=[]
        for i,p in enumerate(people):
            own=[r for r in records if r['persona']==i and (state is None or r['estado']==state) and (not only_load or r['en_carga'])]
            hours=[sum(r['horas'] or 0 for r in own if r['sprint']==sid) for sid in payload['sprintsCSV']]
            base=p['datos_base']
            detail.append([p['nombre'],p['activo'],base.get('Celula interna BD'),base.get('Coordinador'),base.get('Líder'),*hours,sum(hours)])
        totals=[sum(row[j] or 0 for row in detail) for j in range(5,len(headers))]
        rows.append([label,None,None,None,None,*totals]); levels.append(0)
        rows.extend(detail); levels.extend([1]*len(detail))
    for name in ('Matriz asignadas','Matriz cerradas','Matriz pendientes'):
        if name in book: book.remove(book[name])
    source.title='Matriz por estado'
    flat_outline('Matriz por estado',headers,rows,levels).freeze_panes='F2'
    for ws in book:
        ws._charts.clear()
