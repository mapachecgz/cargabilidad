"""Resúmenes limitados a recursos identificados y dashboard local."""
from collections import defaultdict, Counter
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.styles import PatternFill, Font, Alignment


def add_metrics(book, sheet, records, people, sprints, resource_rows, cutoff, alerts):
    def order(s): return (0, int(s)) if s.isdigit() else (1, s)
    linked = [r for r in records if r['persona'] is not None]
    ids = sorted({r['sprint'] for r in linked}, key=order)
    current = [s for s,v in sprints.items() if v['activo'] and v['inicio'] <= cutoff.isoformat() <= v['fin']]
    history = defaultdict(lambda: [0.0]*6)
    all_hours = defaultdict(float)
    for r in linked:
        key = (r['persona'], r['sprint'])
        all_hours[key] += r['horas'] or 0
        if not r['en_carga']: continue
        h = history[key]
        h[0] += r['horas'] or 0
        h[1] += (r['horas'] or 0) if r['estado']=='Closed' else 0
        h[2] += (r['horas'] or 0) if r['pendiente'] else 0
        h[3] += (r['horas'] or 0) if not r['pendiente'] and r['estado']!='Closed' else 0
        h[4] += int(r['horas'] is None)
        h[5] += int(r['pendiente'] and r['horas'] is None)
    caps = {(r[0],r[2]):r for r in resource_rows}
    # Aggregate only periods present in the identified population, never future unused sprints.
    periods = [s for s in ids if s in sprints]
    total_rows, historic_rows = [], []
    for i,p in enumerate(people):
        h = [sum(history[(i,s)][j] for s in ids) for j in range(6)]
        cp = [caps[(p['documento'],s)] for s in periods if (p['documento'],s) in caps]
        available = sum(c[9] or 0 for c in cp)
        missing_sprints = sum(1 for s in ids if s not in sprints and history[(i,s)][0]>0)
        known_pending = sum(history[(i,s)][2] for s in periods)
        ratio = known_pending/available if available else None
        total_rows.append([p['documento'],p['nombre'],p['activo'],h[0],h[1],h[2],h[3],available,ratio,h[5],missing_sprints])
        historic_rows.append([p['documento'],p['nombre'],*h[:4],h[4],sum(all_hours[(i,s)] for s in ids)])
    cw=sheet('Carga recursos',['Documento','Recurso','Activo','Horas asignadas','Horas cerradas','Horas pendientes','Horas otros finales','Disponibles en sprints del CSV','Carga pendiente con fechas','Pendientes sin estimacion','Sprints sin fechas'],total_rows)
    for c in cw['I'][1:]: c.number_format='0.0%'
    sheet('Historico recursos',['Documento','Recurso','Horas asignadas','Horas cerradas','Horas pendientes','Horas otros finales','Elementos sin estimacion','Horas todos los tipos'],historic_rows)
    for label,offset in [('Matriz asignadas',0),('Matriz cerradas',1),('Matriz pendientes',2),('Matriz total',None)]:
        data=[]
        for i,p in enumerate(people):
            base=p['datos_base']
            hours=[all_hours[(i,s)] if offset is None else history[(i,s)][offset] for s in ids]
            data.append([p['nombre'],p['activo'],base.get('Celula interna BD'),base.get('Coordinador'),base.get('Líder'),*hours,sum(hours)])
        matrix=sheet(label,['Recurso','Activo','Celula interna','Coordinador','Lider']+['SP '+(s or 'Sin sprint') for s in ids]+['Total'],data)
        matrix.freeze_panes='F2'
        if ids:
            matrix.conditional_formatting.add(f'F2:{matrix.cell(matrix.max_row,matrix.max_column).coordinate}',ColorScaleRule(start_type='min',start_color='F2F2F0',mid_type='percentile',mid_value=50,mid_color='F6C5CD',end_type='max',end_color='C8102E'))
    sprint_rows, dashboard_data = [], []
    for sid in sorted(set(ids)|set(current), key=order):
        h=[sum(history[(i,sid)][j] for i in range(len(people))) for j in range(6)]
        cp=[caps.get((p['documento'],sid)) for p in people]
        available=sum(c[9] or 0 for c in cp if c) if sid in sprints else None
        excess=sum(bool(c and c[3] and c[9] is not None and c[6]>0 and (c[9]==0 or c[6]>=c[16])) for c in cp)
        no_load=sum(bool(c and c[3] and c[9] and history[(i,sid)][0]==0) for i,c in enumerate(cp))
        inactive=sum(not p['activo'] and history[(i,sid)][2]>0 for i,p in enumerate(people))
        elements=sum(r['sprint']==sid and r['en_carga'] for r in linked)
        sprint_rows.append([sid,elements,h[0],h[1],h[2],h[3],h[4],available,h[2]/available if available else None,h[0]/available if available else None,excess,no_load,inactive,h[1]/h[0] if h[0] else None,h[5], 'Sin registros de carga' if not elements else 'Con registros de carga'])
        for i,p in enumerate(people):
            c=caps.get((p['documento'],sid)); hp=history[(i,sid)]
            base=p['datos_base']
            dashboard_data.append({'sprint':sid,'nombre':p['nombre'],'activo':p['activo'],'celula':base.get('Celula interna BD') or 'Sin celula','horas':hp[0],'cerradas':hp[1],'pendientes':hp[2],'otros':hp[3],'faltantes':hp[5],'disponibles':c[9] if c else None,'estado':c[11] if c else 'SIN FECHAS','total':all_hours[(i,sid)]})
    sw=sheet('Metricas sprint',['Sprint','Elementos de carga','Horas asignadas','Horas cerradas','Horas pendientes','Horas otros finales','Elementos sin estimacion','Horas disponibles','Carga pendiente','Carga asignada','Activos con exceso de pendientes','Activos sin horas registradas','Inactivos con pendientes','Porcentaje horas cerradas','Pendientes sin estimacion','Datos del CSV'],sprint_rows)
    for col in ('I','J','N'):
        for c in sw[col][1:]: c.number_format='0.0%'
    # Totales por EVC y celula, con detalle desplegable por sprint.
    org=defaultdict(lambda:[0,0.0,0.0,0.0,0.0])
    for r in linked:
        b=org[(r['evc'],r['celula'],r['sprint'])]; b[0]+=1
        if r['en_carga']:
            b[1]+=r['horas'] or 0
            b[2]+=(r['horas'] or 0) if r['estado']=='Closed' else 0
            b[3]+=(r['horas'] or 0) if r['pendiente'] else 0
        b[4]+=r['horas'] or 0
    headers=['EVC Azure','Celula Azure','Sprint','Registros Azure','Horas asignadas','Horas cerradas','Horas pendientes','Horas todos los tipos']
    detail_org=[list(k[:2])+[k[2] or 'Sin sprint']+v for k,v in sorted(org.items(),key=lambda item:(item[0][0],item[0][1],order(item[0][2])))]
    sheet('Detalle celulas EVC',headers,detail_org)
    data_org=[]; outlines=[]; summaries=[]
    for evc in sorted({k[0] for k in org}):
        evc_row=len(data_org)+2
        totals=[sum(v[j] for k,v in org.items() if k[0]==evc) for j in range(5)]
        data_org.append([evc,'TOTAL EVC','Todos',*totals]); summaries.append(evc_row)
        for cell in sorted({k[1] for k in org if k[0]==evc}):
            cell_row=len(data_org)+2
            totals=[sum(v[j] for k,v in org.items() if k[:2]==(evc,cell)) for j in range(5)]
            data_org.append([evc,cell,'TOTAL CELULA',*totals]); summaries.append(cell_row)
            for key in sorted((k for k in org if k[:2]==(evc,cell)),key=lambda k:order(k[2])):
                data_org.append([evc,cell,key[2] or 'Sin sprint',*org[key]])
            outlines.append((cell_row+1,len(data_org)+1,2,cell_row))
        outlines.append((evc_row+1,len(data_org)+1,1,evc_row))
    ow=sheet('Celulas y EVC',headers,data_org)
    ow.sheet_properties.outlinePr.summaryBelow=False
    ow._tables.clear(); ow.auto_filter.ref=None
    for start,end,level,summary in sorted(outlines,key=lambda g:g[2]):
        ow.row_dimensions.group(start,end,outline_level=level,hidden=True)
        ow.row_dimensions[summary].collapsed=True
    for row in summaries:
        for c in ow[row]:
            c.font=Font(bold=True,color='8F0B17'); c.fill=PatternFill('solid',fgColor='F2F2F0')
    # Collapsible rows: one visible total per internal team; its resources start collapsed.
    team_data=[]; groups=[]
    team_names=sorted({p['datos_base'].get('Celula interna BD') or 'Sin celula' for p in people})
    for name in team_names:
        members=[i for i,p in enumerate(people) if (p['datos_base'].get('Celula interna BD') or 'Sin celula')==name]
        start=len(team_data)+2
        totals=[sum(total_rows[i][j] for i in members) for j in (3,4,5,7)]
        team_data.append([name,'TOTAL CELULA',len(members),*totals])
        for i in members:
            r=total_rows[i]
            team_data.append([name,r[1],int(r[2]),r[3],r[4],r[5],r[7]])
        groups.append((start,start+1,len(team_data)+1))
    tw=sheet('Equipos internos',['Celula interna','Recurso o subtotal','Recursos / activo','Horas asignadas','Horas cerradas','Horas pendientes','Horas disponibles'],team_data)
    tw.sheet_properties.outlinePr.summaryBelow=False
    for summary,start,end in groups:
        tw.row_dimensions.group(start,end,outline_level=1,hidden=True)
        tw.row_dimensions[summary].collapsed=True
        for c in tw[summary]: c.font=Font(bold=True,color='8F0B17'); c.fill=PatternFill('solid',fgColor='F2F2F0')
    # Do not filter subtotal rows independently from their children.
    tw._tables.clear(); tw.auto_filter.ref=None
    focus=[d for d in dashboard_data if d['sprint'] in current]
    fw=sheet('Sprint vigente',['Sprint','Recurso','Activo','Horas asignadas','Horas cerradas','Horas pendientes','Horas disponibles','Estado pendientes','Pendientes sin estimacion'],[[d[k] for k in ('sprint','nombre','activo','horas','cerradas','pendientes','disponibles','estado','faltantes')] for d in focus])
    resource_ids={r['id'] for r in linked}
    quality=Counter(a[0] for a in alerts if a[1] in resource_ids or a[0]=='SOBRECARGA')
    sheet('Calidad datos',['Tipo alerta','Eventos de recursos identificados'],list(quality.most_common()))
    # Compact dashboard, with explanations in a separate dictionary.
    active=sum(p['activo'] for p in people)
    indicators=[['Fecha de corte',cutoff.isoformat()],['Recursos configurados',len(people)],['Recursos activos',active],['Recursos inactivos',len(people)-active],['Elementos de recursos identificados',len(linked)],['Horas asignadas (tipos de carga)',sum(r[3] for r in total_rows)],['Horas cerradas',sum(r[4] for r in total_rows)],['Horas pendientes',sum(r[5] for r in total_rows)],['Sprints vigentes',', '.join(current) or 'Ninguno']]
    for sid in current:
        r=next(r for r in sprint_rows if r[0]==sid)
        indicators.extend([[f'Sprint {sid}: datos disponibles',r[15]],[f'Sprint {sid}: horas registradas',r[2]],[f'Sprint {sid}: horas de trabajo pendientes',r[4]],[f'Sprint {sid}: horas disponibles del equipo',r[7]],[f'Sprint {sid}: recursos con exceso de pendientes',r[10]],[f'Sprint {sid}: recursos activos sin horas registradas',r[11]],[f'Sprint {sid}: inactivos con trabajo pendiente',r[12]],[f'Sprint {sid}: pendientes sin estimacion',r[14]]])
    dashboard=sheet('Tablero',['Indicador','Valor'],indicators)
    dashboard.column_dimensions['A'].width=54; dashboard.column_dimensions['B'].width=34
    dashboard.sheet_view.showGridLines=False
    if sprint_rows:
        chart=LineChart(); chart.title='Recursos identificados: horas por sprint'; chart.y_axis.title='Horas'
        chart.add_data(Reference(sw,min_col=3,max_col=5,min_row=1,max_row=sw.max_row),titles_from_data=True)
        chart.set_categories(Reference(sw,min_col=1,min_row=2,max_row=sw.max_row)); chart.width=24; chart.height=11
        dashboard.add_chart(chart,'D2')
    if focus:
        chart=BarChart(); chart.type='bar'; chart.title='Sprint vigente: pendientes y disponibilidad'
        chart.add_data(Reference(fw,min_col=6,max_col=7,min_row=1,max_row=fw.max_row),titles_from_data=True)
        chart.set_categories(Reference(fw,min_col=2,min_row=2,max_row=fw.max_row)); chart.width=24; chart.height=20
        dashboard.add_chart(chart,'D25')
    glossary=[['Recursos identificados','Personas de recursos.json vinculadas por correo o nombre; incluye activos e inactivos.'],['Dedicacion','Porcentaje de jornada para este equipo: 100% usa la jornada completa; 50% usa la mitad. Reduce horas brutas y descuento festivo; actividades fijas por sprint.'],['Horas asignadas','Story Points de Task y HA configurados, en cualquier estado.'],['Horas cerradas','Horas estimadas de elementos Closed. No son horas ejecutadas.'],['Horas otros finales','Horas de estados finales configurados distintos de Closed, como Removed y Resolved.'],['Matriz total','Story Points de TODOS los tipos y estados; incluye Issue y Epic si tienen horas. No equivale a la carga Task/HA.'],['Horas disponibles','Dias habiles del sprint por jornada y dedicacion, menos festivos, ausencias y actividades.'],['Carga recursos','Una fila por persona. Disponibilidad sumada en los sprints con registros identificados en el CSV y fechas conocidas. Ratio usa pendientes de esos mismos sprints; se muestra cantidad de sprints sin fechas.'],['Historico recursos','Una fila por persona: totales historicos, sin dimensiones organizativas. Para comparar sprints use matrices.'],['Sin registros de carga','El CSV no trae elementos de carga para el sprint. Cero horas conocidas no garantiza ausencia de trabajo.'],['Festivos','Horas por festivo multiplicadas por el numero de fechas festivas habiles dentro del sprint, sin repetir fechas.'],['Disponibilidad historica','Se aplica la configuracion actual de activo, dedicacion y ausencias; no existe historial de cambios.'],['Registros Azure','Numero de elementos (IDs) del CSV vinculados a recursos identificados, de todos los tipos y estados. No son personas ni horas.'],['Celulas y EVC','Use + para desplegar EVC, celulas y sprints. No sume subtotales y detalle. Detalle celulas EVC permite filtrar por EVC, celula y sprint con listas desplegables.'],['Equipos internos','Use los botones + de Excel para desplegar los recursos de cada celula.'],['Actualizacion','Coloque un solo CSV en entrada y regenere. Las configuraciones se leen exclusivamente de config.']]
    sheet('Guia de lectura',['Concepto','Descripcion'],glossary)
    book._sheets.remove(dashboard); book._sheets.insert(0,dashboard)
    if cw.max_row>1: cw.conditional_formatting.add(f'I2:I{cw.max_row}',CellIsRule(operator='greaterThanOrEqual',formula=['1'],fill=PatternFill('solid',fgColor='F6C5CD')))
    return {'fecha':cutoff.isoformat(),'actuales':current,'sprintsCSV':ids,'sprints':sorted(set(ids)|set(current),key=order),'recursos':dashboard_data,'indicadores':indicators}
